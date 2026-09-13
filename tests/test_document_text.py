"""Тесты извлечения TXT и DOCX без ElevenLabs."""

from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from services.document_text import (
    DocumentEmptyError,
    DocumentError,
    extract_text_from_document,
    extract_text_from_docx,
    extract_text_from_txt,
)


def _write(path: Path, data: bytes) -> Path:
    path.write_bytes(data)
    return path


class TxtExtractionTests(unittest.TestCase):
    def test_utf8_is_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = _write(Path(tmp) / "a.txt", "Привет, мир!".encode("utf-8"))
            self.assertEqual(extract_text_from_txt(path), "Привет, мир!")

    def test_utf8_bom_is_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = _write(Path(tmp) / "a.txt", "Инструкция".encode("utf-8-sig"))
            self.assertEqual(extract_text_from_txt(path), "Инструкция")

    def test_cp1251_is_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = _write(Path(tmp) / "a.txt", "Текст Windows".encode("cp1251"))
            self.assertEqual(extract_text_from_txt(path), "Текст Windows")

    def test_empty_txt_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = _write(Path(tmp) / "a.txt", b"")
            with self.assertRaises(DocumentEmptyError):
                extract_text_from_txt(path)

    def test_binary_txt_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = _write(Path(tmp) / "a.txt", bytes(range(256)) + b"\x00\x01\x02")
            with self.assertRaises(DocumentError):
                extract_text_from_txt(path)


class DocxExtractionTests(unittest.TestCase):
    def test_paragraphs_are_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.docx"
            document = Document()
            document.add_paragraph("Первый абзац.")
            document.add_paragraph("Второй абзац.")
            document.save(path)
            text = extract_text_from_docx(path)
            self.assertIn("Первый абзац.", text)
            self.assertIn("Второй абзац.", text)

    def test_headings_are_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.docx"
            document = Document()
            document.add_heading("Заголовок Венеции", level=1)
            document.add_paragraph("Обычный текст.")
            document.save(path)
            text = extract_text_from_docx(path)
            self.assertIn("Заголовок Венеции", text)

    def test_tables_are_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.docx"
            document = Document()
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Север"
            table.cell(0, 1).text = "Юг"
            table.cell(1, 0).text = "Запад"
            table.cell(1, 1).text = "Восток"
            document.save(path)
            text = extract_text_from_docx(path)
            self.assertIn("Север", text)
            self.assertIn("Восток", text)

    def test_paragraph_and_table_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.docx"
            document = Document()
            document.add_paragraph("До таблицы")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Ячейка"
            document.add_paragraph("После таблицы")
            document.save(path)
            text = extract_text_from_docx(path)
            self.assertLess(text.find("До таблицы"), text.find("Ячейка"))
            self.assertLess(text.find("Ячейка"), text.find("После таблицы"))

    def test_empty_docx_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.docx"
            Document().save(path)
            with self.assertRaises(DocumentEmptyError):
                extract_text_from_docx(path)

    def test_corrupted_docx_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = _write(Path(tmp) / "a.docx", b"PK\x03\x04not-a-document")
            with self.assertRaises(DocumentError):
                extract_text_from_docx(path)

    def test_fake_zip_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.docx"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("readme.txt", "not docx")
            with self.assertRaises(DocumentError):
                extract_text_from_docx(path)

    def test_document_wrapper_uses_extension(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = _write(Path(tmp) / "note.txt", "Hello".encode("utf-8"))
            self.assertEqual(extract_text_from_document(path, "txt"), "Hello")


if __name__ == "__main__":
    unittest.main()
